"""Чтение базы планирования из .xlsx в нормализованные Point (раздел 4, 10 ТЗ)."""
from __future__ import annotations

from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from optimizer.models import Point
from optimizer.normalizer import canonical_address, normalize_street

from .columns import resolve_column
from .errors import ValidationError
from .validator import validate_planning_header


def _optional_float(value) -> float | None:
    """Читает опциональное число. Пусто/не число -> None (не учитываем)."""
    if value is None:
        return None
    s = str(value).strip().replace("%", "").replace(",", ".").strip()
    if s == "":
        return None
    try:
        return float(s)
    except ValueError:
        return None


def _display_value(value, number_format: str | None) -> str:
    """Строка, как её видит пользователь в Excel (учитывает процентный формат).

    openpyxl в data_only=True отдаёт «сырое» число (0.9), теряя формат ячейки
    «0%», из-за чего в выгрузке появлялись цифры вместо процентов. Здесь мы
    восстанавливаем видимое значение, не меняя данные: 0.9 в формате «0%» →
    «90%», строки — как есть, пусто — пусто.
    """
    if value is None:
        return ""
    if isinstance(value, str):
        return value.strip()
    if isinstance(value, bool):
        return str(value)
    if number_format and "%" in number_format and isinstance(value, (int, float)):
        return f"{value * 100:g}%"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _lookup_ci(raw: dict, name: str):
    """Поиск значения по колонке без учёта регистра (Оценка/результат/Результат)."""
    for k, v in raw.items():
        if str(k).strip().lower() == name.lower():
            return v
    return None


def _read_header_and_rows(fileobj):
    try:
        wb = load_workbook(fileobj, read_only=True, data_only=True)
    except (BadZipFile, InvalidFileException, OSError) as exc:
        raise ValidationError(
            "Файл не является корректным .xlsx. Проверьте, что загружаете именно "
            "Excel-файл (.xlsx), а не служебный файл (~$…) или файл другого формата.",
            {"reason": f"{type(exc).__name__}: {exc}"},
        )
    ws = wb.active
    # values_only=False — нужен доступ к number_format для восстановления «%».
    rows = ws.iter_rows(values_only=False)
    header_row = next(rows, None)
    if header_row is None:
        return None, rows
    header = [c.value for c in header_row]
    return header, rows


def parse_planning(fileobj) -> list[Point]:
    """Парсит базу планирования. Первая строка — заголовок.

    Поднимает ValidationError с понятным сообщением при структурных ошибках.
    """
    header, rows = _read_header_and_rows(fileobj)
    if header is None:
        raise ValidationError("Файл пуст")
    header = [str(c).strip() if c is not None else "" for c in header]
    validate_planning_header(header)

    points: list[Point] = []
    seen_codes: set[str] = set()

    for row_num, cells in enumerate(rows, start=2):
        if cells is None or all(c.value is None or str(c.value).strip() == "" for c in cells):
            continue
        raw: dict[str, object] = {}
        original: dict[str, object] = {}
        for i, c in enumerate(cells):
            if i >= len(header):
                break
            key = header[i]
            val = c.value
            raw[key] = "" if val is None else val
            if key.strip().lower() in ("оценка", "результат"):
                original[key] = _display_value(val, c.number_format)
            else:
                original[key] = "" if val is None else val

        # Единый код не обязателен — генерируем уникальный, если пусто.
        code = str(resolve_column(raw, "Единый код")).strip()
        if not code:
            code = f"ТОЧКА-{row_num}"
        if code in seen_codes:
            raise ValidationError(f"Дублирующийся Единый код: {code}")
        seen_codes.add(code)

        # Цикличность не обязательна — по умолчанию 1.
        freq_str = str(resolve_column(raw, "Цикличность")).strip()
        if freq_str == "":
            frequency = 1
        else:
            try:
                frequency = int(float(freq_str))
            except ValueError:
                raise ValidationError(
                    f"Некорректная цикличность в строке {row_num}: «{freq_str}»"
                )
        if frequency < 1:
            raise ValidationError(f"Цикличность должна быть >= 1 (строка {row_num})")

        locality = str(resolve_column(raw, "Населенный пункт")).strip()
        metro = str(resolve_column(raw, "Станция метро")).strip()
        house = str(resolve_column(raw, "Номер дома")).strip()
        street_name = str(resolve_column(raw, "Название улицы")).strip()
        street = normalize_street(
            f"{resolve_column(raw, 'Тип улицы')} {street_name}".strip()
        )
        address = canonical_address(locality, metro, street, house)

        # Нет полного адреса (город или улица пусты) -> «неопределённая точка».
        incomplete = locality == "" or street_name == ""

        points.append(
            Point(
                id=code,
                trade_rep_code=str(resolve_column(raw, "Код торгового представителя")).strip(),
                frequency=frequency,
                locality=locality,
                street=street,
                house=house,
                metro=metro,
                normalized_address=address,
                score=_optional_float(_lookup_ci(raw, "Оценка")),
                result=_optional_float(_lookup_ci(raw, "Результат")),
                original=original,
                incomplete=incomplete,
                row=row_num,
            )
        )

    if not points:
        raise ValidationError("База не содержит точек")
    return points


def read_template_header(fileobj) -> list[str]:
    """Возвращает столбцы шаблона маршрутного листа (первая строка)."""
    header, _ = _read_header_and_rows(fileobj)
    if header is None:
        raise ValidationError("Шаблон пуст")
    return [str(c).strip() if c is not None else "" for c in header]
