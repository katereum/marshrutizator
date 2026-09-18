import io
import re
import unittest
from datetime import date, datetime

from openpyxl import load_workbook, Workbook

from exporter.exporter import export_route
from optimizer.models import Point
from optimizer.pipeline import build_route
from parser.columns import RESULT_COLUMNS


def template_bytes():
    wb = Workbook()
    ws = wb.active
    ws.append(RESULT_COLUMNS)
    ws.append(["" for _ in RESULT_COLUMNS])  # строка-образец формата
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class TestExporter(unittest.TestCase):
    def test_export_preserves_header_and_fills_values(self):
        points = [
            Point(
                id="A001",
                trade_rep_code="TP1",
                frequency=1,
                locality="Москва",
                street="ул ленина",
                house="1",
                original={
                    "Единый код": "A001",
                    "Старый код": "1001",
                    "Населенный пункт": "Москва",
                    "Тип улицы": "ул.",
                    "Название улицы": "Ленина",
                    "Номер дома": "1",
                    "Код торгового представителя": "TP1",
                },
            )
        ]
        result = build_route(points, date(2026, 9, 1), date(2026, 9, 30))
        data = export_route(io.BytesIO(template_bytes()), result, {p.id: p for p in points})

        wb = load_workbook(io.BytesIO(data))
        ws = wb.active
        headers = [ws.cell(row=1, column=c).value for c in range(1, len(RESULT_COLUMNS) + 1)]
        self.assertEqual(headers, RESULT_COLUMNS)

        found = None
        for r in range(2, ws.max_row + 1):
            if ws.cell(row=r, column=4).value == "A001":  # колонка «Единый код»
                found = r
                break
        self.assertIsNotNone(found)
        self.assertEqual(ws.cell(row=found, column=1).value, 1)  # «№ п/п»
        date_cell = ws.cell(row=found, column=2)
        self.assertIsInstance(date_cell.value, (date, datetime))  # настоящая дата
        self.assertEqual(date_cell.number_format, "DD.MM.YYYY")  # 01.07 = 1 июля
        self.assertIn(ws.cell(row=found, column=3).value, [
            "Понедельник", "Вторник", "Среда", "Четверг", "Пятница", "Суббота", "Воскресенье",
        ])
        self.assertEqual(ws.cell(row=found, column=4).value, "A001")

    def test_export_fills_optional_score_and_result(self):
        from parser.columns import OPTIONAL_RESULT_COLUMNS

        wb = Workbook()
        ws = wb.active
        ws.append(RESULT_COLUMNS + OPTIONAL_RESULT_COLUMNS)
        ws.append(["" for _ in RESULT_COLUMNS + OPTIONAL_RESULT_COLUMNS])
        buf = io.BytesIO()
        wb.save(buf)

        points = [
            Point(
                id="A001", trade_rep_code="TP1", frequency=1, locality="Москва",
                street="ул ленина", house="1", score=3.0, result=77.5,
                original={
                    "Единый код": "A001",
                    "Код торгового представителя": "TP1",
                    "Оценка": 3,
                    "Результат": "77.5%",
                },
            )
        ]
        result = build_route(points, date(2026, 9, 1), date(2026, 9, 30))
        data = export_route(io.BytesIO(buf.getvalue()), result, {p.id: p for p in points})

        out = load_workbook(io.BytesIO(data))
        sheet = out.active
        header = [sheet.cell(row=1, column=c).value for c in range(1, sheet.max_column + 1)]
        oi = header.index("Оценка") + 1
        ri = header.index("Результат") + 1
        self.assertEqual(sheet.cell(row=2, column=oi).value, "3")
        self.assertEqual(sheet.cell(row=2, column=ri).value, "77.5%")  # «%» сохранён

    def test_export_adds_summary_sheet(self):
        from optimizer.priority import FOCUS_FIX

        points = [
            Point(id="A001", trade_rep_code="TP1", frequency=1, locality="Москва",
                  street="ул ленина", house="1", score=1.0, result=10.0,
                  original={"Единый код": "A001", "Код торгового представителя": "TP1"}),
            Point(id="A002", trade_rep_code="TP1", frequency=1, locality="Москва",
                  street="ул ленина", house="2", score=5.0, result=100.0,
                  original={"Единый код": "A002", "Код торгового представителя": "TP1"}),
        ]
        result = build_route(points, date(2026, 9, 1), date(2026, 9, 30), focus=FOCUS_FIX)
        data = export_route(io.BytesIO(template_bytes()), result, {p.id: p for p in points})

        wb = load_workbook(io.BytesIO(data))
        self.assertIn("Сводка", wb.sheetnames)
        ws = wb["Сводка"]
        self.assertEqual(ws["B1"].value, "Сначала проблемные")
        self.assertEqual(ws["B2"].value, "2 из 2")  # обе с оценкой
        self.assertEqual(ws["B3"].value, "2 из 2")  # обе с результатом

    def test_export_auto_adds_score_and_result_columns(self):
        # Шаблон БЕЗ колонок Оценка/Результат — они должны добавиться сами.
        points = [
            Point(
                id="A001", trade_rep_code="TP1", frequency=1, locality="Москва",
                street="ул ленина", house="1", score=3.0, result=80.0,
                original={
                    "Единый код": "A001",
                    "Код торгового представителя": "TP1",
                    "Оценка": 3,
                    "Результат": "80%",
                },
            )
        ]
        result = build_route(points, date(2026, 9, 1), date(2026, 9, 30))
        data = export_route(io.BytesIO(template_bytes()), result, {p.id: p for p in points})

        wb = load_workbook(io.BytesIO(data))
        ws = wb.active
        header = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
        self.assertIn("Оценка", header)
        self.assertIn("Результат", header)
        oi = header.index("Оценка") + 1
        ri = header.index("Результат") + 1
        self.assertEqual(ws.cell(row=2, column=oi).value, "3")
        self.assertEqual(ws.cell(row=2, column=ri).value, "80%")  # «%» сохранён

    def test_export_case_insensitive_result_column(self):
        # Шаблон с «результат» (маленькая буква) — заполняем его, без дубля «Результат».
        wb = Workbook()
        ws = wb.active
        ws.append(RESULT_COLUMNS + ["Оценка", "результат"])
        ws.append(["" for _ in RESULT_COLUMNS + ["Оценка", "результат"]])
        buf = io.BytesIO()
        wb.save(buf)

        points = [
            Point(
                id="A001", trade_rep_code="TP1", frequency=1, locality="Москва",
                street="ул ленина", house="1", score=3.0, result=80.0,
                original={
                    "Единый код": "A001",
                    "Код торгового представителя": "TP1",
                    "Оценка": 3,
                    "результат": "80%",
                },
            )
        ]
        result = build_route(points, date(2026, 9, 1), date(2026, 9, 30))
        data = export_route(io.BytesIO(buf.getvalue()), result, {p.id: p for p in points})

        out = load_workbook(io.BytesIO(data))
        sheet = out.active
        header = [sheet.cell(row=1, column=c).value for c in range(1, sheet.max_column + 1)]
        self.assertIn("результат", header)
        self.assertNotIn("Результат", header)  # не создаём дубль
        ri = header.index("результат") + 1
        self.assertEqual(sheet.cell(row=2, column=ri).value, "80%")

    def test_result_display_adds_percent_to_number(self):
        # Результат числом (80) в базе -> в выгрузке «80%».
        points = [
            Point(
                id="A001", trade_rep_code="TP1", frequency=1, locality="Москва",
                street="ул ленина", house="1", score=3.0, result=80.0,
                original={"Единый код": "A001", "Код торгового представителя": "TP1", "Результат": 80},
            )
        ]
        result = build_route(points, date(2026, 9, 1), date(2026, 9, 30))
        data = export_route(io.BytesIO(template_bytes()), result, {p.id: p for p in points})
        wb = load_workbook(io.BytesIO(data))
        ws = wb.active
        header = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
        ri = header.index("Результат") + 1
        self.assertEqual(ws.cell(row=2, column=ri).value, "80%")


if __name__ == "__main__":
    unittest.main()
