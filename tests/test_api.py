import io
import unittest
from unittest.mock import patch

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from backend.app import _build_road_matrix, _human_address, app, store
from optimizer.models import Point
from parser.columns import PLANNING_COLUMNS, RESULT_COLUMNS

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def make_xlsx(header, rows):
    wb = Workbook()
    ws = wb.active
    ws.append(header)
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def planning_row(code, cycle=1):
    return [code, "1", "Актив", "Москва", "ул.", "Ленина", "1", "", "", "Тверская",
            "", "П", "С", "К", "SUP", "TP1", cycle]


class TestAPI(unittest.TestCase):
    def setUp(self):
        store._lock.acquire()
        try:
            store._files.clear()
            store._jobs.clear()
        finally:
            store._lock.release()
        self.client = TestClient(app)

    def test_end_to_end(self):
        planning = make_xlsx(
            PLANNING_COLUMNS,
            [planning_row("A001", 1), planning_row("A002", 1), planning_row("A003", 1)],
        )
        r = self.client.post(
            "/api/upload/planning",
            files={"file": ("planning.xlsx", planning, XLSX_MIME)},
        )
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.json()["row_count"], 3)
        pid = r.json()["file_id"]

        template = make_xlsx(RESULT_COLUMNS, [["" for _ in RESULT_COLUMNS]])
        r = self.client.post(
            "/api/upload/route-template",
            files={"file": ("template.xlsx", template, XLSX_MIME)},
        )
        self.assertEqual(r.status_code, 200)
        tid = r.json()["file_id"]

        with patch("backend.app.geocoder") as mock_geo, patch("backend.app.road_distance") as mock_road, patch("backend.app.route_geometry") as mock_route_geo:
            mock_geo.geocode.return_value = (55.75, 37.61)
            mock_road.matrix.return_value = [
                [0.0, 1.0, 1.0],
                [1.0, 0.0, 1.0],
                [1.0, 1.0, 0.0],
            ]
            mock_route_geo.route.return_value = None
            r = self.client.post(
                "/api/optimize",
                json={
                    "planning_file_id": pid,
                    "route_template_file_id": tid,
                    "period_start": "2026-09-01",
                    "period_end": "2026-09-30",
                },
            )
        self.assertEqual(r.status_code, 202)
        jid = r.json()["job_id"]

        r = self.client.get(f"/api/result/{jid}")
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.json()["status"], "COMPLETED")
        self.assertEqual(r.json()["stats"][0]["total_visits"], 3)
        self.assertEqual(r.json()["stats"][0]["violations"], [])

        r = self.client.get(f"/api/result/{jid}/geojson")
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.json()["type"], "FeatureCollection")

        r = self.client.get(f"/api/download/{jid}")
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.headers["content-type"], XLSX_MIME)

    def test_optimize_falls_back_without_road_matrix(self):
        planning = make_xlsx(
            PLANNING_COLUMNS,
            [planning_row("A001", 1), planning_row("A002", 1)],
        )
        r = self.client.post(
            "/api/upload/planning",
            files={"file": ("planning.xlsx", planning, XLSX_MIME)},
        )
        pid = r.json()["file_id"]

        template = make_xlsx(RESULT_COLUMNS, [["" for _ in RESULT_COLUMNS]])
        r = self.client.post(
            "/api/upload/route-template",
            files={"file": ("template.xlsx", template, XLSX_MIME)},
        )
        tid = r.json()["file_id"]

        with patch("backend.app.geocoder") as mock_geo, patch("backend.app.road_distance") as mock_road, patch("backend.app.route_geometry") as mock_route_geo:
            mock_geo.geocode.return_value = (55.75, 37.61)
            mock_road.matrix.return_value = None
            mock_route_geo.route.return_value = None
            r = self.client.post(
                "/api/optimize",
                json={
                    "planning_file_id": pid,
                    "route_template_file_id": tid,
                    "period_start": "2026-09-01",
                    "period_end": "2026-09-30",
                },
            )

        # Нет дорожной матрицы -> откат на прямые расстояния с предупреждением.
        self.assertEqual(r.status_code, 202)
        jid = r.json()["job_id"]
        r2 = self.client.get(f"/api/result/{jid}")
        self.assertEqual(r2.json()["status"], "COMPLETED")
        self.assertTrue(any("прямым" in w for w in r2.json()["stats"][0]["warnings"]))

    def test_optimize_capacity_mismatch_needs_confirmation(self):
        # 3 визита, но лимит 2/день на 22 рабочих дня = 44 слота — не сходится.
        planning = make_xlsx(
            PLANNING_COLUMNS,
            [planning_row("A001", 1), planning_row("A002", 1), planning_row("A003", 1)],
        )
        r = self.client.post("/api/upload/planning", files={"file": ("planning.xlsx", planning, XLSX_MIME)})
        pid = r.json()["file_id"]

        template = make_xlsx(RESULT_COLUMNS, [["" for _ in RESULT_COLUMNS]])
        r = self.client.post("/api/upload/route-template", files={"file": ("template.xlsx", template, XLSX_MIME)})
        tid = r.json()["file_id"]

        body = {
            "planning_file_id": pid,
            "route_template_file_id": tid,
            "period_start": "2026-09-01",
            "period_end": "2026-09-30",
            "points_per_day": 2,
        }
        with patch("backend.app.geocoder") as mock_geo, patch("backend.app.road_distance") as mock_road, patch("backend.app.route_geometry") as mock_route_geo:
            mock_geo.geocode.return_value = (55.75, 37.61)
            mock_road.matrix.return_value = [[0.0] * 3 for _ in range(3)]
            mock_route_geo.route.return_value = None
            r = self.client.post("/api/optimize", json=body)

        self.assertEqual(r.status_code, 200)
        data = r.json()
        self.assertTrue(data["needs_confirmation"])
        self.assertEqual(data["total_visits"], 3)
        self.assertLess(data["difference"], 0)  # не хватает

        # confirm=true → оптимизация проходит.
        body["confirm"] = True
        with patch("backend.app.geocoder") as mock_geo, patch("backend.app.road_distance") as mock_road, patch("backend.app.route_geometry") as mock_route_geo:
            mock_geo.geocode.return_value = (55.75, 37.61)
            mock_road.matrix.return_value = [[0.0] * 3 for _ in range(3)]
            mock_route_geo.route.return_value = None
            r = self.client.post("/api/optimize", json=body)

        self.assertEqual(r.status_code, 202)
        self.assertIn("job_id", r.json())

    def test_optimize_address_confirmation_flow(self):
        # Одна точка с полным адресом, одна — без города («неопределённая»).
        planning = make_xlsx(
            PLANNING_COLUMNS,
            [
                planning_row("A001", 1),
                ["A002", "1", "Актив", "", "ул.", "Ленина", "2", "", "", "", "", "П", "С", "К", "SUP", "TP1", 1],
            ],
        )
        r = self.client.post("/api/upload/planning", files={"file": ("planning.xlsx", planning, XLSX_MIME)})
        pid = r.json()["file_id"]

        body = {
            "planning_file_id": pid,
            "period_start": "2026-09-01",
            "period_end": "2026-09-30",
        }
        with patch("backend.app.geocoder") as mock_geo:
            mock_geo.geocode.return_value = (55.75, 37.61)
            r = self.client.post("/api/optimize", json=body)

        self.assertEqual(r.status_code, 200)
        data = r.json()
        self.assertTrue(data["needs_address_confirmation"])
        self.assertEqual(data["incomplete_count"], 1)

        # Подтверждаем — маршрут по полной точке, неполная уходит в отдельный лист.
        body["confirm_address"] = True
        with patch("backend.app.geocoder") as mock_geo:
            mock_geo.geocode.return_value = (55.75, 37.61)
            r = self.client.post("/api/optimize", json=body)

        self.assertEqual(r.status_code, 202)
        jid = r.json()["job_id"]
        r = self.client.get(f"/api/download/{jid}")
        self.assertEqual(r.status_code, 200)
        wb = load_workbook(io.BytesIO(r.content))
        self.assertIn("Неопределенные точки", wb.sheetnames)

    def test_optimize_duplicate_confirmation_flow(self):
        # Две строки с одним кодом — предупреждаем, по подтверждению считаем один раз.
        planning = make_xlsx(PLANNING_COLUMNS, [planning_row("A001", 1), planning_row("A001", 1)])
        r = self.client.post("/api/upload/planning", files={"file": ("planning.xlsx", planning, XLSX_MIME)})
        pid = r.json()["file_id"]

        body = {
            "planning_file_id": pid,
            "period_start": "2026-09-01",
            "period_end": "2026-09-30",
        }
        r = self.client.post("/api/optimize", json=body)
        self.assertEqual(r.status_code, 200)
        data = r.json()
        self.assertTrue(data["needs_duplicate_confirmation"])
        self.assertEqual(data["duplicate_count"], 1)

        body["confirm_duplicates"] = True
        with patch("backend.app.geocoder") as mock_geo:
            mock_geo.geocode.return_value = (55.75, 37.61)
            r = self.client.post("/api/optimize", json=body)
        self.assertEqual(r.status_code, 202)

    def test_human_address_resolves_column_aliases(self):
        # Колонки названы иначе («Город», «Улица», «Дом», «Тип») — адрес должен собираться по синонимам.
        p = Point(
            id="A1", trade_rep_code="TP1", frequency=1,
            locality="Москва", street="шоссе Бесединское", house="15",
            original={"Город": "Москва", "Тип": "шоссе", "Улица": "Бесединское", "Дом": "15"},
        )
        self.assertEqual(_human_address(p), "Москва, шоссе Бесединское, 15")

    def test_optimize_with_aliased_address_columns(self):
        header = ["Код точки", "Город", "Тип", "Улица", "Дом", "Сколько раз посещаем в месяц", "ТП"]
        rows = [["A001", "Москва", "шоссе", "Бесединское", "15", 1, "TP1"]]
        r = self.client.post("/api/upload/planning", files={"file": ("p.xlsx", make_xlsx(header, rows), XLSX_MIME)})
        self.assertEqual(r.status_code, 200, r.text)
        pid = r.json()["file_id"]

        body = {"planning_file_id": pid, "period_start": "2026-09-01", "period_end": "2026-09-30"}
        with patch("backend.app.geocoder") as mock_geo:
            mock_geo.geocode.return_value = (55.75, 37.61)
            r = self.client.post("/api/optimize", json=body)
        self.assertEqual(r.status_code, 202, r.text)
        mock_geo.geocode.assert_called()
        addr = mock_geo.geocode.call_args[0][0]
        self.assertIn("Москва", addr)
        self.assertIn("Бесединское", addr)

    def test_upload_missing_column_returns_error(self):
        # «Название улицы» (адрес) — обязательна в шапке.
        header = ["Единый код", "Населенный пункт", "Цикличность"]
        rows = [["A001", "Москва", 1]]
        r = self.client.post(
            "/api/upload/planning",
            files={"file": ("p.xlsx", make_xlsx(header, rows), XLSX_MIME)},
        )
        self.assertEqual(r.status_code, 400)
        self.assertEqual(r.json()["error"]["code"], "VALIDATION_ERROR")

    def test_optimize_without_template(self):
        # Без маршрутного листа-шаблона — только база планирования.
        planning = make_xlsx(
            PLANNING_COLUMNS,
            [planning_row("A001", 1), planning_row("A002", 1)],
        )
        r = self.client.post("/api/upload/planning", files={"file": ("planning.xlsx", planning, XLSX_MIME)})
        self.assertEqual(r.status_code, 200)
        pid = r.json()["file_id"]

        with patch("backend.app.geocoder") as mock_geo, patch("backend.app.road_distance") as mock_road, patch("backend.app.route_geometry") as mock_route_geo:
            mock_geo.geocode.return_value = (55.75, 37.61)
            mock_road.matrix.return_value = [[0.0, 1.0], [1.0, 0.0]]
            mock_route_geo.route.return_value = None
            r = self.client.post(
                "/api/optimize",
                json={
                    "planning_file_id": pid,
                    "period_start": "2026-09-01",
                    "period_end": "2026-09-30",
                },
            )

        self.assertEqual(r.status_code, 202)
        jid = r.json()["job_id"]

        r = self.client.get(f"/api/download/{jid}")
        self.assertEqual(r.status_code, 200)
        wb = load_workbook(io.BytesIO(r.content))
        ws = wb.active
        header = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
        self.assertIn("Дата визита", header)
        self.assertIn("Единый код", header)
        # Оценка/Приоритет появляются только если есть в базе (здесь их нет).
        self.assertNotIn("Оценка", header)
        self.assertNotIn("Приоритет", header)

    def test_upload_missing_secondary_columns_ok(self):
        # Второстепенные колонки («Код Супервайзера», «Старый код», «Статус» …) не обязательны.
        header = ["Единый код", "Населенный пункт", "Название улицы", "Цикличность"]
        rows = [["A001", "Москва", "Ленина", 1]]
        r = self.client.post(
            "/api/upload/planning",
            files={"file": ("p.xlsx", make_xlsx(header, rows), XLSX_MIME)},
        )
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.json()["row_count"], 1)

    def test_upload_alias_column_names_ok(self):
        # Колонки могут называться иначе: «Код точки» = «Единый код», «Метро» = «Станция метро».
        header = ["Код точки", "Населенный пункт", "Название улицы", "Цикличность", "Код СВ", "Метро"]
        rows = [["A001", "Москва", "Ленина", 1, "SUP1", "Тверская"]]
        r = self.client.post(
            "/api/upload/planning",
            files={"file": ("p.xlsx", make_xlsx(header, rows), XLSX_MIME)},
        )
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.json()["row_count"], 1)

    def test_upload_wrong_extension(self):
        r = self.client.post(
            "/api/upload/planning",
            files={"file": ("p.csv", b"not excel", "text/csv")},
        )
        self.assertEqual(r.status_code, 400)
        self.assertEqual(r.json()["error"]["code"], "FILE_ERROR")


class TestBuildRoadMatrix(unittest.TestCase):
    def test_chunks_requests_by_road_matrix_chunk(self):
        points = [
            Point(id=f"p{i}", trade_rep_code="T", frequency=1,
                  latitude=55.0 + i * 0.01, longitude=37.0 + i * 0.01)
            for i in range(3)
        ]
        calls = []

        def fake_matrix(origins, destinations):
            calls.append((len(origins), len(destinations)))
            return [[0.0 if o == d else 1.0 for d in destinations] for o in origins]

        with patch.dict("os.environ", {"ROAD_MATRIX_CHUNK": "2"}), \
                patch("backend.app.road_distance") as mock_road:
            mock_road.matrix.side_effect = fake_matrix
            matrix = _build_road_matrix(points, None)

        # 3 точки при размере блока 2 → блоки [2, 1], итого 4 запроса.
        self.assertEqual(calls, [(2, 2), (2, 1), (1, 2), (1, 1)])
        self.assertEqual(len(matrix), 3 * 3)
        self.assertEqual(matrix[("p0", "p2")], 1.0)


if __name__ == "__main__":
    unittest.main()
