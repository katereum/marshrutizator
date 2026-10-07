"""Этап 9 — полный end-to-end и проверка критериев приёмки (specification.md §12)."""
import io
import unittest
import zipfile
from collections import Counter
from datetime import datetime
from unittest.mock import patch

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from backend.app import app, store
from parser.columns import RESULT_COLUMNS

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

# Колонки, значения которых переносятся из базы без искажений (критерий 14).
POINT_COLUMNS = [
    "Старый код", "Статус", "Населенный пункт", "Тип улицы", "Название улицы",
    "Номер дома", "Номер строения", "Номер корпуса", "Станция метро",
    "Дополнительное описание месторасположения точки", "Партнер", "Субдилер",
    "Субканал", "Код Супервайзера", "Код торгового представителя",
]


def read_original(path):
    wb = load_workbook(path)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    header = [str(c).strip() if c is not None else "" for c in rows[0]]
    out = {}
    for row in rows[1:]:
        d = {header[i]: (row[i] if i < len(row) and row[i] is not None else "") for i in range(len(header))}
        code = str(d.get("Единый код", "")).strip()
        if code:
            out[code] = d
    return out


def read_output_rows(content):
    wb = load_workbook(io.BytesIO(content))
    ws = wb.active
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[3] is None or str(row[3]).strip() == "":
            continue
        rows.append(row)
    return rows


class TestAcceptance(unittest.TestCase):
    def setUp(self):
        store._lock.acquire()
        try:
            store._files.clear()
            store._jobs.clear()
        finally:
            store._lock.release()
        self.client = TestClient(app)

    def _upload(self, path, endpoint):
        with open(path, "rb") as f:
            content = f.read()
        r = self.client.post(endpoint, files={"file": (path.split("/")[-1], content, XLSX_MIME)})
        self.assertEqual(r.status_code, 200)
        return r.json()["file_id"]

    def test_full_pipeline_meets_acceptance_criteria(self):
        # Критерии 1–3: загрузка обоих файлов, проверка структуры, чтение данных.
        pid = self._upload("sample_data/planning.xlsx", "/api/upload/planning")
        tid = self._upload("sample_data/route_template.xlsx", "/api/upload/route-template")

        # Разные адреса -> разные координаты (иначе все точки «в одном месте»,
        # и новый этап сборки по месту схлопнет их в один день).
        _seen: dict[str, int] = {}
        _counter = [0]

        def _geo(address):
            if address not in _seen:
                _seen[address] = _counter[0]
                _counter[0] += 1
            i = _seen[address]
            return (55.70 + (i % 10) * 0.01, 37.40 + (i // 10) * 0.01)

        with patch("backend.app.geocoder") as mock_geo, patch("backend.app.road_distance") as mock_road, patch("backend.app.route_geometry") as mock_route_geo:
            mock_geo.geocode.side_effect = _geo
            mock_road.matrix.side_effect = lambda origins, destinations: [
                [0.0 if o == d else 1.0 for d in destinations]
                for o in origins
            ]
            mock_route_geo.route.return_value = None
            r = self.client.post(
                "/api/optimize",
                json={
                    "planning_file_id": pid,
                    "route_template_file_id": tid,
                    "period_start": "2026-09-01",
                    "period_end": "2026-09-30",
                },
            )
        self.assertEqual(r.status_code, 202)
        jid = r.json()["job_id"]

        r = self.client.get(f"/api/result/{jid}")
        self.assertEqual(r.status_code, 200)
        body = r.json()
        self.assertEqual(body["status"], "COMPLETED")
        stats = body["stats"]
        self.assertEqual(len(stats), 2)  # два сотрудника

        # Критерий 12: результат прошёл автопроверку.
        for s in stats:
            self.assertEqual(s["violations"], [])

        # Критерий 9: нагрузка равномерна (разброс ≤ 1).
        for s in stats:
            self.assertLessEqual(max(s["per_day_load"]) - min(s["per_day_load"]), 1)

        # Критерии 13, 15: формируется и скачивается корректный Excel.
        r = self.client.get(f"/api/download/{jid}")
        self.assertEqual(r.status_code, 200)
        if r.headers["content-type"] == "application/zip":
            files = {}
            with zipfile.ZipFile(io.BytesIO(r.content)) as zf:
                for name in zf.namelist():
                    files[name] = zf.read(name)
        else:
            files = {"route.xlsx": r.content}
        self.assertGreaterEqual(len(files), 2)

        original = read_original("sample_data/planning.xlsx")

        # Критерий 4: все точки базы учтены.
        all_codes = {str(row[3]).strip() for content in files.values() for row in read_output_rows(content)}
        self.assertEqual(all_codes, set(original.keys()))

        # Критерии 5–7: цикличность точная, нет лишних/пропущенных посещений.
        freq_count = Counter(str(row[3]).strip() for content in files.values() for row in read_output_rows(content))
        for code, d in original.items():
            self.assertEqual(freq_count.get(code, 0), int(d["Цикличность"]), f"цикличность {code}")

        # Критерий 8: визиты только на рабочие дни.
        for content in files.values():
            for row in read_output_rows(content):
                raw_date = row[1]
                if isinstance(raw_date, datetime):
                    d = raw_date.date()
                else:
                    d = datetime.strptime(str(raw_date), "%d.%m.%Y").date()
                self.assertLess(d.weekday(), 5, f"выходной день {row[1]}")

        # Критерий 14: исходные данные точек не искажены.
        col_index = {c: i for i, c in enumerate(RESULT_COLUMNS)}
        for content in files.values():
            for row in read_output_rows(content):
                code = str(row[3]).strip()
                orig = original[code]
                for col in POINT_COLUMNS:
                    expected = str(orig.get(col, "")).strip()
                    actual = str(row[col_index[col]] if row[col_index[col]] is not None else "").strip()
                    self.assertEqual(actual, expected, f"{code}: {col}")


if __name__ == "__main__":
    unittest.main()
