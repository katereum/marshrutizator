"""Заполнение шаблона маршрутного листа или генерация листа с нуля (раздел 24 ТЗ, R6)."""
from __future__ import annotations

import io
from copy import copy
from typing import BinaryIO

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from optimizer.models import Point, RouteResult
from parser.columns import RESULT_COLUMNS

WEEKDAYS_RU = [
    "Понедельник", "Вторник", "Среда", "Четверг", "Пятница", "Суббота", "Воскресенье",
]

FOCUS_LABELS = {
    "economy": "Экономия бензина",
    "fix_problems": "Сначала проблемные",
    "top_performers": "Сначала лучшие",
}

# Колонки готового маршрутного листа, когда шаблон не загружен.
DEFAULT_RESULT_COLUMNS = list(RESULT_COLUMNS) + ["Оценка", "Результат", "Приоритет"]


def _raw(point: Point, key: str):
    """Сырое значение из базы (сохраняет «%»), пустое -> прочерк. Без регистра."""
    o = point.original or {}
    for k, v in o.items():
        if str(k).strip().lower() == key.lower():
            if v is None:
                return "—"
            s = str(v).strip()
            return s if s != "" else "—"
    return "—"


def _result_display(point: Point) -> str:
    """Результат — как в исходной базе, без изменений.

    Вводные данные не преобразуем: процент уже восстановлен парсером по формату
    ячейки («0%» → «90%»), поэтому здесь просто возвращаем исходное значение.
    """
    return _raw(point, "Результат")


def _find_column(col_index: dict, name: str):
    """Колонка по имени без учёта регистра (Оценка/результат/Результат)."""
    for hdr, col in col_index.items():
        if str(hdr).strip().lower() == name.lower():
            return col
    return None


def _find_header(ws) -> tuple[int, dict[str, int]] | None:
    for r in range(1, min(ws.max_row, 10) + 1):
        vals = [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
        names = {str(v).strip(): c for c, v in enumerate(vals, start=1) if v is not None and str(v).strip()}
        if "Единый код" in names and "Дата визита" in names:
            return r, names
    return None


def _copy_style(dst, src) -> None:
    dst.font = copy(src.font)
    dst.border = copy(src.border)
    dst.fill = copy(src.fill)
    dst.alignment = copy(src.alignment)
    dst.number_format = src.number_format
    dst.protection = copy(src.protection)


def _add_summary_sheet(wb, result: RouteResult, points_by_id: dict[str, Point]) -> None:
    """Добавляет лист «Сводка»: акцент, учёт Оценки/Результата и приоритеты."""
    ws = wb.create_sheet("Сводка")

    # Точки именно этого маршрута (в порядке появления), без дублей.
    ids: list[str] = []
    for day in result.days:
        for vid in day.ordered_visits:
            pid = vid.rsplit("#", 1)[0]
            if pid not in ids:
                ids.append(pid)
    points = [points_by_id[pid] for pid in ids if pid in points_by_id]

    n_score = sum(1 for p in points if p.score is not None)
    n_result = sum(1 for p in points if p.result is not None)

    ws["A1"] = "Акцент маршрута"
    ws["B1"] = FOCUS_LABELS.get(result.focus, result.focus)
    ws["A2"] = "Точек с Оценкой"
    ws["B2"] = f"{n_score} из {len(points)}"
    ws["A3"] = "Точек с Результатом"
    ws["B3"] = f"{n_result} из {len(points)}"

    # Предупреждения (если есть — например, откат на прямые расстояния).
    header_row = 5
    if result.warnings:
        ws.cell(row=header_row, column=1, value="Предупреждения")
        r = header_row
        for w in result.warnings:
            r += 1
            ws.cell(row=r, column=1, value="• " + w)
        header_row = r + 2

    ws.cell(row=header_row, column=1, value="Код")
    ws.cell(row=header_row, column=2, value="Оценка")
    ws.cell(row=header_row, column=3, value="Результат")
    ws.cell(row=header_row, column=4, value="Приоритет")

    ordered = sorted(points, key=lambda p: result.priority.get(p.id, 0.0), reverse=True)
    for i, p in enumerate(ordered):
        row = header_row + 1 + i
        ws.cell(row=row, column=1, value=p.id)
        ws.cell(row=row, column=2, value=_raw(p, "Оценка"))
        ws.cell(row=row, column=3, value=_result_display(p))
        ws.cell(row=row, column=4, value=round(result.priority.get(p.id, 0.0), 2))


def _route_row_values(point: Point, day, num: int, result: RouteResult) -> dict:
    orig = point.original
    return {
        "№ п/п": num,
        "Дата визита": day.date,
        "День недели": WEEKDAYS_RU[day.weekday],
        "Единый код": orig.get("Единый код", point.id),
        "Старый код": orig.get("Старый код", ""),
        "Статус": orig.get("Статус", ""),
        "Населенный пункт": orig.get("Населенный пункт", ""),
        "Тип улицы": orig.get("Тип улицы", ""),
        "Название улицы": orig.get("Название улицы", ""),
        "Номер дома": orig.get("Номер дома", ""),
        "Номер строения": orig.get("Номер строения", ""),
        "Номер корпуса": orig.get("Номер корпуса", ""),
        "Станция метро": orig.get("Станция метро", ""),
        "Дополнительное описание месторасположения точки": orig.get(
            "Дополнительное описание месторасположения точки", ""
        ),
        "Партнер": orig.get("Партнер", ""),
        "Субдилер": orig.get("Субдилер", ""),
        "Субканал": orig.get("Субканал", ""),
        "Код Супервайзера": orig.get("Код Супервайзера", ""),
        "Код торгового представителя": orig.get("Код торгового представителя", ""),
        "Сколько раз посещаем в месяц": point.frequency,
        "Оценка": _raw(point, "Оценка"),
        "Результат": _result_display(point),
        "Приоритет": "" if result.priority.get(point.id) is None else round(result.priority[point.id], 2),
    }


def _fill_route_rows(ws, col_index: dict, write_row: int, sample_row, result: RouteResult, points_by_id: dict) -> None:
    """Заполняет строки маршрута по колонкам col_index (имя -> индекс)."""
    for day in result.days:
        num = 1  # нумерация «№ п/п» начинается заново каждый день
        for vid in day.ordered_visits:
            point = points_by_id.get(vid.rsplit("#", 1)[0])
            if point is None:
                continue
            values = _route_row_values(point, day, num, result)
            for name, value in values.items():
                column = _find_column(col_index, name)
                if column is None:
                    continue
                cell = ws.cell(row=write_row, column=column, value=value)
                if sample_row is not None:
                    _copy_style(cell, ws.cell(row=sample_row, column=column))
                if name == "Дата визита":
                    # Настоящая дата, всегда в формате ДД.ММ.ГГГГ (1 июля = 01.07).
                    cell.number_format = "DD.MM.YYYY"
            num += 1
            write_row += 1


def _export_route_from_scratch(result: RouteResult, points_by_id: dict[str, Point]) -> bytes:
    """Создаёт маршрутный лист с нуля: стандартные колонки + Оценка/Результат/Приоритет."""
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Маршрут"
    header_row = 1
    col_index = {}
    for c, name in enumerate(DEFAULT_RESULT_COLUMNS, start=1):
        cell = ws.cell(row=header_row, column=c, value=name)
        cell.font = Font(bold=True)
        col_index[name] = c
    _fill_route_rows(ws, col_index, header_row + 1, None, result, points_by_id)
    _add_summary_sheet(wb, result, points_by_id)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def export_route(template: BinaryIO | None, result: RouteResult, points_by_id: dict[str, Point]) -> bytes:
    """Заполняет шаблон (или создаёт лист с нуля) и возвращает .xlsx в виде bytes."""
    if template is None:
        return _export_route_from_scratch(result, points_by_id)

    wb = load_workbook(template)
    ws = wb.active

    found = _find_header(ws)
    if found is None:
        raise ValueError("В шаблоне не найдена строка заголовка (Единый код + Дата визита)")
    header_row, col_index = found

    # Автодобавляем опциональные колонки, если их нет в шаблоне (без регистра):
    # так Оценка/Результат из базы всегда попадают в выгрузку без правки шаблона.
    next_col = ws.max_column + 1
    existing = {str(k).strip().lower() for k in col_index}
    for name in ("Оценка", "Результат", "Приоритет"):
        if name.lower() not in existing:
            col_index[name] = next_col
            ws.cell(row=header_row, column=next_col, value=name)
            next_col += 1

    sample_row = header_row + 1
    has_sample = ws.cell(row=sample_row, column=col_index["Единый код"]).value not in (None, "")

    write_row = sample_row if has_sample else header_row + 1
    _fill_route_rows(ws, col_index, write_row, sample_row if has_sample else None, result, points_by_id)

    _add_summary_sheet(wb, result, points_by_id)

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
